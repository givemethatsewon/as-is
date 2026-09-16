from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import csv
from pathlib import Path
from typing import Any

import pandas as pd
from fastapi import UploadFile
from openpyxl import load_workbook


class ParseError(ValueError):
    pass


SHEET_HINTS = {
    "imports": ("원상태진행", "원상태 진행", "재고", "수입", "stock", "import"),
    "exports": ("수출", "수출 양식", "invoice", "export"),
}
HEADER_HINTS = {
    "imports": (
        "수입신고번호",
        "신고일자",
        "수리일",
        "원산지",
        "세번",
        "hs code",
        "란번호",
        "행번호",
        "part number",
        "판매부번",
        "규격",
        "수량",
        "잔량",
    ),
    "exports": (
        "order no",
        "수출신고번호",
        "신고일자",
        "part number",
        "규격1",
        "규격2",
        "description",
        "u/price",
        "ready to ship qty",
        "ready to ship",
        "qty",
        "amount",
        "수입신고번호",
        "수리일",
        "원산지",
        "세번",
        "사용수량",
    ),
}
EXPORT_DATE_HEADERS = {
    "export_date",
    "수출일",
    "수출일자",
    "수출예정일",
    "신고일자",
    "shipping date",
    "invoice date",
}
SUBHEADER_HINTS = {"qty", "amount", "수량", "금액"}


async def read_upload_rows(file: UploadFile, upload_type: str | None = None) -> list[dict[str, Any]]:
    filename = file.filename or ""
    content = await file.read()
    if not content:
        raise ParseError("Uploaded file is empty.")

    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        frame = pd.read_csv(StringIO(text), dtype=str).fillna("")
    elif filename.lower().endswith((".xlsx", ".xlsm")):
        frame = _read_excel_table(content, upload_type)
    else:
        raise ParseError("Only CSV, XLSX, and XLSM files are supported.")

    frame.columns = [str(column).strip() for column in frame.columns]
    return frame.to_dict(orient="records")


def iter_file_rows(path: str | Path, upload_type: str | None = None):
    """Yield data rows without loading the entire workbook into memory."""
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as stream:
            yield from csv.DictReader(stream)
        return
    if suffix not in {".xlsx", ".xlsm"}:
        raise ParseError("Only CSV, XLSX, and XLSM files are supported.")

    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=False, keep_links=False)
    try:
        sheet, header_index, sampled_rows = _select_streaming_sheet(workbook, upload_type)
        headers, data_offset = _streaming_headers(sampled_rows, header_index)
        keep = [(index, header) for index, header in enumerate(headers) if header]
        for values in sheet.iter_rows(min_row=data_offset + 1, values_only=True):
            row = {header: values[index] if index < len(values) else None for index, header in keep}
            if any(value not in (None, "") for value in row.values()):
                yield row
    finally:
        workbook.close()


def _select_streaming_sheet(workbook, upload_type: str | None):
    candidates = []
    hints = HEADER_HINTS.get(upload_type or "", HEADER_HINTS["imports"] + HEADER_HINTS["exports"])
    for sheet_index, sheet in enumerate(workbook.worksheets):
        sampled = [list(row) for row in sheet.iter_rows(min_row=1, max_row=20, values_only=True)]
        if not sampled:
            continue
        best_header = 0
        best_score = -1
        for index, values in enumerate(sampled):
            cells = [_normalize_header(value) for value in values]
            score = sum(_cell_matches_any_hint(cell, hints) for cell in cells if cell)
            if score > best_score:
                best_header = index
                best_score = score
        candidates.append((_sheet_score(sheet.title, upload_type) + best_score, -sheet_index, sheet, best_header, sampled))
    if not candidates:
        raise ParseError("Excel file has no readable rows.")
    candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
    _, _, sheet, header_index, sampled = candidates[0]
    return sheet, header_index, sampled


def _streaming_headers(rows: list[list[Any]], header_index: int) -> tuple[list[str], int]:
    primary = [_clean_header(value) for value in rows[header_index]]
    if header_index + 1 >= len(rows):
        return [_dedupe_header(header, primary[:index]) for index, header in enumerate(primary)], header_index + 1
    secondary = [_clean_header(value) for value in rows[header_index + 1]]
    if not _looks_like_subheader_row(secondary):
        return primary, header_index + 1
    headers = []
    last_primary = ""
    for main, sub in zip(primary, secondary, strict=False):
        if main:
            last_primary = main
        if main and sub:
            headers.append(f"{main} {sub}")
        elif main:
            headers.append(main)
        elif sub:
            headers.append(sub if not last_primary else sub)
        else:
            headers.append("")
    return [_dedupe_header(header, headers[:index]) for index, header in enumerate(headers)], header_index + 2


def _read_excel_table(content: bytes, upload_type: str | None) -> pd.DataFrame:
    workbook = pd.ExcelFile(BytesIO(content))
    candidates: list[tuple[int, int, str, pd.DataFrame]] = []

    for sheet_index, sheet_name in enumerate(workbook.sheet_names):
        raw = pd.read_excel(workbook, sheet_name=sheet_name, header=None, dtype=str).fillna("")
        if raw.empty:
            continue
        header_row, header_score = _best_header_row(raw, upload_type)
        sheet_score = _sheet_score(sheet_name, upload_type)
        candidates.append((sheet_score + header_score, -sheet_index, sheet_name, _table_from_header(raw, header_row)))

    if not candidates:
        raise ParseError("Excel file has no readable rows.")

    candidates.sort(key=lambda candidate: (candidate[0], candidate[1]), reverse=True)
    frame = candidates[0][3]
    if upload_type == "exports":
        frame = _ensure_export_date(frame)
    return frame


def _best_header_row(frame: pd.DataFrame, upload_type: str | None) -> tuple[int, int]:
    hints = HEADER_HINTS.get(upload_type or "", HEADER_HINTS["imports"] + HEADER_HINTS["exports"])
    best_index = 0
    best_score = -1
    max_scan_rows = min(len(frame.index), 20)

    for index in range(max_scan_rows):
        cells = [_normalize_header(value) for value in frame.iloc[index].tolist()]
        score = sum(_cell_matches_any_hint(cell, hints) for cell in cells if cell)
        if score > best_score:
            best_index = index
            best_score = score
    return best_index, best_score


def _table_from_header(frame: pd.DataFrame, header_row: int) -> pd.DataFrame:
    raw_headers, data_start = _headers_from_row(frame, header_row)
    data = frame.iloc[data_start:].copy()
    keep_indices = [index for index, header in enumerate(raw_headers) if header]
    headers = [_dedupe_header(raw_headers[index], raw_headers[:index]) for index in keep_indices]

    if not keep_indices:
        raise ParseError("Excel sheet has no recognizable header row.")

    data = data.iloc[:, keep_indices]
    data.columns = headers
    data = data.replace("", pd.NA).dropna(how="all").fillna("")
    return data


def _headers_from_row(frame: pd.DataFrame, header_row: int) -> tuple[list[str], int]:
    primary = [_clean_header(value) for value in frame.iloc[header_row].tolist()]
    if header_row + 1 >= len(frame.index):
        return primary, header_row + 1

    secondary = [_clean_header(value) for value in frame.iloc[header_row + 1].tolist()]
    if not _looks_like_subheader_row(secondary):
        return primary, header_row + 1

    headers: list[str] = []
    last_primary = ""
    for main, sub in zip(primary, secondary, strict=False):
        if main:
            last_primary = main
        if main and sub:
            headers.append(f"{main} {sub}")
        elif main:
            headers.append(main)
        elif sub:
            headers.append(sub if not last_primary else sub)
        else:
            headers.append("")
    return headers, header_row + 2


def _looks_like_subheader_row(cells: list[str]) -> bool:
    return any(_normalize_header(cell) in SUBHEADER_HINTS for cell in cells if cell)


def _sheet_score(sheet_name: str, upload_type: str | None) -> int:
    hints = SHEET_HINTS.get(upload_type or "", ())
    normalized_sheet_name = _normalize_header(sheet_name)
    for priority, hint in enumerate(hints):
        if _normalize_header(hint) == normalized_sheet_name:
            return 100 - priority
    return 0


def _ensure_export_date(frame: pd.DataFrame) -> pd.DataFrame:
    has_export_date = any(_normalize_header(column) in EXPORT_DATE_HEADERS for column in frame.columns)
    if has_export_date:
        return frame
    frame = frame.copy()
    frame["export_date"] = date.today().isoformat()
    return frame


def _cell_matches_any_hint(cell: str, hints: tuple[str, ...]) -> int:
    return int(any(cell == _normalize_header(hint) or _normalize_header(hint) in cell for hint in hints))


def _clean_header(value: Any) -> str:
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def _normalize_header(value: Any) -> str:
    return " ".join(_clean_header(value).casefold().split())


def _dedupe_header(header: str, previous_headers: list[str]) -> str:
    if header not in previous_headers:
        return header
    count = sum(1 for previous in previous_headers if previous == header) + 1
    return f"{header}_{count}"


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def optional_text(value: Any) -> str | None:
    text = clean_text(value)
    return text or None


def parse_date(value: Any, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        raise ValueError(f"{field} is required.")
    for date_format in ("%Y-%m-%d", "%Y%m%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    parsed = pd.to_datetime(text, errors="coerce")
    if pd.isna(parsed):
        raise ValueError(f"{field} must be a valid date.")
    return parsed.date()


def parse_positive_int(value: Any, field: str) -> int:
    text = clean_text(value).replace(",", "")
    if not text:
        raise ValueError(f"{field} is required.")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{field} must be numeric.") from exc
    if number <= 0 or number != number.to_integral_value():
        raise ValueError(f"{field} must be a positive integer.")
    return int(number)


def parse_non_negative_int(value: Any, field: str) -> int:
    text = clean_text(value).replace(",", "")
    if not text:
        raise ValueError(f"{field} is required.")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{field} must be numeric.") from exc
    if number < 0 or number != number.to_integral_value():
        raise ValueError(f"{field} must be a non-negative integer.")
    return int(number)


def parse_decimal(value: Any, field: str) -> Decimal | None:
    text = clean_text(value).replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{field} must be numeric.") from exc
