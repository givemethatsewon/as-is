from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.services.reports import matching_run_workbook
from app.services.uploads import confirm_batch, confirm_match_run, preview_export_run, preview_imports


def test_matching_run_workbook_repeats_export_fields_and_adds_inventory_sheet(db_session) -> None:
    imports = [
        {
            "import_declaration_no": "IMP-A",
            "import_accepted_date": "2026-01-01",
            "origin": "CN",
            "hs_code": "8501",
            "line_no": "1",
            "row_no": "1",
            "part_number": "PN-1",
            "spec": "IMPORT-SPEC-A",
            "import_qty": "4",
            "qty_unit": "PC",
        },
        {
            "import_declaration_no": "IMP-B",
            "import_accepted_date": "2026-01-02",
            "origin": "CN",
            "hs_code": "8502",
            "line_no": "1",
            "row_no": "2",
            "part_number": "PN-1",
            "spec": "IMPORT-SPEC-B",
            "import_qty": "3",
            "qty_unit": "PC",
        },
    ]
    import_preview = preview_imports(db_session, imports, "imports.xlsx")
    confirm_batch(db_session, import_preview.batch.id)
    exports = [
        {
            "export_date": "2026-02-01",
            "order_no": "ORDER-9",
            "seq_no": "7",
            "origin": "CN",
            "part_number": "PN-1",
            "hs_code": "EXPORT-HS-IGNORED",
            "description": "EXPORT-DESCRIPTION",
            "unit_price": "2.5",
            "required_qty": "9",
            "amount": "22.5",
        }
    ]
    run = preview_export_run(db_session, exports, "exports.xlsx", eligibility_days=720)
    confirm_match_run(db_session, run.batch.id)

    workbook = load_workbook(BytesIO(matching_run_workbook(db_session, run.batch.id)))

    assert workbook.sheetnames == ["수출 결과", "원상태잔량"]
    result = workbook["수출 결과"]
    headers = [cell.value for cell in result[1]]
    values = [dict(zip(headers, row, strict=True)) for row in result.iter_rows(min_row=2, values_only=True)]
    assert [row["Order No"] for row in values] == ["ORDER-9", "ORDER-9", "ORDER-9"]
    assert [row["Seq No"] for row in values] == ["7", "7", "7"]
    assert [row["U/Price"] for row in values] == [2.5, 2.5, 2.5]
    assert [row["Amount"] for row in values] == [22.5, 22.5, 22.5]
    assert [row["수입신고번호"] for row in values] == ["IMP-A", "IMP-B", "NO MATCH"]
    assert [row["세번"] for row in values[:2]] == ["8501", "8502"]
    assert [row["수입 규격"] for row in values[:2]] == ["IMPORT-SPEC-A", "IMPORT-SPEC-B"]
    assert values[-1]["부족 수량"] == 2
    assert values[-1]["매칭 상태"] == "NO MATCH"

    inventory = workbook["원상태잔량"]
    inventory_headers = [cell.value for cell in inventory[1]]
    inventory_rows = [
        dict(zip(inventory_headers, row, strict=True)) for row in inventory.iter_rows(min_row=2, values_only=True)
    ]
    assert [row["수입신고번호"] for row in inventory_rows] == ["IMP-A", "IMP-B"]
    assert [row["잔량"] for row in inventory_rows] == [0, 0]


def test_original_upload_can_be_downloaded_from_authenticated_endpoint(client, db_session, tmp_path) -> None:
    original = tmp_path / "original.xlsx"
    original.write_bytes(b"original-file")
    from app.models import UploadBatch

    batch = UploadBatch(
        upload_type="imports",
        filename="수입 원본.xlsx",
        source_path=str(original),
        source_sha256="hash",
        source_size_bytes=13,
    )
    db_session.add(batch)
    db_session.commit()

    response = client.get(f"/api/upload-batches/{batch.id}/original")

    assert response.status_code == 200
    assert response.content == b"original-file"
    assert "filename" in response.headers["content-disposition"]

